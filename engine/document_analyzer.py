import os
import sys
import base64
import tempfile
from pathlib import Path
from engine.config import GEMINI_API_KEY

def extract_text_from_file(file_path: str) -> tuple[str, dict]:
    """Extract raw text and metadata from PDF, DOCX, TXT, CSV, MD, Code files."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = path.suffix.lower()
    text = ""
    metadata = {
        "filename": path.name,
        "extension": ext,
        "size_kb": round(path.stat().st_size / 1024, 1),
        "pages": None
    }

    # 1. PDF Documents
    if ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(str(path))
            metadata["pages"] = len(reader.pages)
            extracted = []
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                extracted.append(f"--- Page {i+1} ---\n{page_text}")
            text = "\n\n".join(extracted)
        except Exception as pe:
            raise RuntimeError(f"Error reading PDF {path.name}: {pe}")

    # 2. Word Documents (.docx)
    elif ext in [".docx", ".doc"]:
        try:
            import docx
            doc = docx.Document(str(path))
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_text:
                        paras.append(row_text)
            text = "\n\n".join(paras)
        except Exception as de:
            raise RuntimeError(f"Error reading Word document {path.name}: {de}")

    # 3. Text, Markdown, CSV, Code, JSON files
    else:
        encodings = ["utf-8", "latin-1", "cp1252"]
        read_success = False
        for enc in encodings:
            try:
                with open(path, "r", encoding=enc, errors="replace") as f:
                    text = f.read()
                read_success = True
                break
            except Exception:
                continue
        if not read_success:
            raise RuntimeError(f"Could not read content from {path.name}")

    text = text.strip()
    words = len(text.split())
    metadata["word_count"] = words
    return text, metadata

def resolve_document_path(query_or_name: str) -> str:
    """Find document file by name or path across standard Windows directories."""
    q = query_or_name.strip().strip('"').strip("'")
    p = Path(q)
    if p.exists() and p.is_file():
        return str(p.resolve())

    user_home = Path(os.environ.get("USERPROFILE", ""))
    search_dirs = [
        Path(r"s:\Project\AI Assistant"),
        user_home / "Desktop",
        user_home / "Documents",
        user_home / "Downloads",
        Path(r"C:\Users\Public\Desktop")
    ]

    clean_target = q.lower()
    # Direct match in directories
    for folder in search_dirs:
        if folder.exists():
            direct = folder / q
            if direct.exists() and direct.is_file():
                return str(direct.resolve())

    # Fuzzy / extension matching in directories
    for folder in search_dirs:
        if folder.exists():
            try:
                for item in folder.iterdir():
                    if item.is_file():
                        name_lower = item.name.lower()
                        stem_lower = item.stem.lower()
                        if clean_target == name_lower or clean_target == stem_lower or clean_target in name_lower:
                            return str(item.resolve())
            except Exception:
                continue

    return ""

def analyze_document(file_path: str, user_question: str = None) -> dict:
    """
    Extract text and analyze document using Gemini.
    Returns structured summary and concise spoken briefing.
    """
    try:
        raw_text, metadata = extract_text_from_file(file_path)
    except Exception as e:
        return {
            "success": False,
            "error": str(e),
            "filename": Path(file_path).name,
            "summary_text": f"Error: {e}",
            "spoken_summary": f"I was unable to read the document {Path(file_path).name}."
        }

    if not raw_text:
        return {
            "success": False,
            "error": "Document is empty",
            "filename": metadata["filename"],
            "summary_text": "The document contains no readable text.",
            "spoken_summary": f"The document {metadata['filename']} appears to be empty."
        }

    # Bound text to first 45,000 characters to ensure ultra-fast generation
    truncated_text = raw_text[:45000]

    prompt = (
        "You are an executive AI Virtual Assistant analyzing a document for your user.\n"
        f"Document: {metadata['filename']} (Size: {metadata['size_kb']} KB, Words: {metadata['word_count']})\n\n"
        f"--- DOCUMENT CONTENT ---\n{truncated_text}\n--- END OF CONTENT ---\n\n"
    )

    if user_question:
        prompt += (
            f"User's Question: {user_question}\n"
            "Please answer the question accurately based on the document.\n"
            "Format your answer with clear key points and end with a section:\n"
            "SPOKEN_SUMMARY: <A concise 1 to 2 sentence answer suitable for voice speech without markdown>"
        )
    else:
        prompt += (
            "Please provide an executive summary of this document with the following structure:\n"
            "### 📌 Executive Overview\n(2-3 sentences summarizing the purpose and core message)\n\n"
            "### 🔑 Key Highlights & Takeaways\n(4-6 bullet points covering main findings, data, or decisions)\n\n"
            "### 💡 Actionable Insights\n(2-3 practical recommendations or next steps)\n\n"
            "SPOKEN_SUMMARY: <A crisp 1 to 2 sentence overview of what the document is about suitable for voice speech without markdown>"
        )

    from engine.llm import get_gemini_client
    from google.genai import types

    client = get_gemini_client()
    if not client:
        return {
            "success": False,
            "error": "Gemini client unavailable",
            "filename": metadata["filename"],
            "summary_text": "AI service currently unreachable.",
            "spoken_summary": "I read the document, but the AI analysis service is temporarily offline."
        }

    summary_text = ""
    candidate_models = ["gemini-3.5-flash-lite", "gemini-3.5-flash", "gemini-3.7-flash"]
    for m in candidate_models:
        try:
            resp = client.models.generate_content(
                model=m,
                contents=[types.Content(role="user", parts=[types.Part.from_text(text=prompt)])],
                config=types.GenerateContentConfig(temperature=0.4)
            )
            summary_text = resp.text.strip()
            break
        except Exception as e:
            print(f"Doc analysis notice ({m}): {e}")
            continue

    if not summary_text:
        return {
            "success": False,
            "error": "Analysis failed",
            "filename": metadata["filename"],
            "summary_text": "Could not generate analysis.",
            "spoken_summary": f"I processed {metadata['filename']}, but could not generate a summary."
        }

    # Extract SPOKEN_SUMMARY
    spoken_summary = ""
    if "SPOKEN_SUMMARY:" in summary_text:
        parts = summary_text.split("SPOKEN_SUMMARY:")
        full_display_summary = parts[0].strip()
        spoken_summary = parts[1].strip()
    else:
        full_display_summary = summary_text
        spoken_summary = f"Here is the executive summary for {metadata['filename']}. I have displayed the key highlights and insights on your screen."

    return {
        "success": True,
        "filename": metadata["filename"],
        "path": file_path,
        "size_kb": metadata["size_kb"],
        "word_count": metadata["word_count"],
        "summary_text": full_display_summary,
        "spoken_summary": spoken_summary
    }

def analyze_document_base64(filename: str, b64_data: str, user_question: str = None) -> dict:
    """Handle document uploaded directly via UI file picker."""
    try:
        # Strip header if present (e.g. data:application/pdf;base64,...)
        if "," in b64_data:
            b64_data = b64_data.split(",", 1)[1]
        decoded = base64.b64decode(b64_data)
        
        # Save to temp file
        ext = os.path.splitext(filename)[1]
        temp_dir = tempfile.gettempdir()
        temp_path = os.path.join(temp_dir, f"dracarys_upload_{filename}")
        with open(temp_path, "wb") as f:
            f.write(decoded)

        result = analyze_document(temp_path, user_question)
        result["filename"] = filename
        return result
    except Exception as e:
        return {
            "success": False,
            "error": str(e),
            "filename": filename,
            "summary_text": f"Error analyzing uploaded file: {e}",
            "spoken_summary": f"There was an error reading the uploaded file {filename}."
        }
